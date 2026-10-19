"""Importadores restringidos, con ubicación de evidencia por fila/página/XML."""
import csv
import hashlib
import io
import re
import zipfile
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from defusedxml import ElementTree as SafeET

from .db import roots

MAX_SIZE = 20 * 1024 * 1024
MAX_UNPACKED = 80 * 1024 * 1024
MAX_ROWS = 50000
FORMATS = {'.xml': 'xml', '.csv': 'csv', '.xlsx': 'xlsx', '.pdf': 'pdf'}
CBC = '{urn:oasis:names:specification:ubl:schema:xsd:CommonBasicComponents-2}'
CAC = '{urn:oasis:names:specification:ubl:schema:xsd:CommonAggregateComponents-2}'
FIELDS = {
    'issuer_ruc': ('ruc_emisor', 'ruc_proveedor'),
    'receiver_ruc': ('ruc_receptor', 'ruc_cliente'),
    'voucher_id': ('serie_numero', 'numero_comprobante', 'comprobante'),
    'total': ('total', 'importe_total'),
    'currency': ('moneda', 'codigo_moneda'),
    'issue_date': ('fecha_emision', 'fecha'),
    'voucher_type': ('tipo_comprobante', 'tipo'),
}
REQUIRED = ('issuer_ruc', 'voucher_id', 'total', 'currency', 'issue_date')


class DocumentError(ValueError):
    pass


def _normalized(text):
    return re.sub(r'\s+', '_', str(text or '').strip().lower().replace('-', '_'))


def _amount(raw):
    s = str(raw).strip()
    if not s or ',' in s or not re.fullmatch(r'[+-]?\d+(?:\.\d{1,2})?', s):
        raise DocumentError(f'Importe no canónico: {raw!r}. Usar punto decimal, sin separador de miles.')
    try:
        d = Decimal(s)
    except InvalidOperation as exc:
        raise DocumentError('Importe no válido') from exc
    if not d.is_finite() or d < 0:
        raise DocumentError('Importe negativo o no finito')
    return str(d.quantize(Decimal('0.01')))


def _date(raw):
    if isinstance(raw, datetime):
        return raw.date().isoformat()
    if isinstance(raw, date):
        return raw.isoformat()
    value = str(raw or '').strip()
    try:
        return date.fromisoformat(value).isoformat()
    except ValueError as exc:
        raise DocumentError(f'Fecha inválida (YYYY-MM-DD): {value}') from exc


def _record(mapping, source_kind, location, ruc):
    issuer = str(mapping.get('issuer_ruc') or '').strip()
    receiver = str(mapping.get('receiver_ruc') or '').strip()
    voucher = str(mapping.get('voucher_id') or '').strip().upper()
    currency = str(mapping.get('currency') or '').strip().upper()
    typ = str(mapping.get('voucher_type') or '01').strip().zfill(2)
    if len(issuer) != 11 or not issuer.isascii() or not issuer.isdigit():
        raise DocumentError(f'{location}: RUC emisor incorrecto')
    if receiver and (len(receiver) != 11 or not receiver.isascii() or not receiver.isdigit()):
        raise DocumentError(f'{location}: RUC receptor incorrecto')
    if not re.fullmatch(r'[A-Z0-9]{1,12}-[A-Z0-9]{1,12}', voucher):
        raise DocumentError(f'{location}: Serie-número no reconocida')
    if not re.fullmatch('[A-Z]{3}', currency):
        raise DocumentError(f'{location}: Moneda inválida')
    if ruc not in (issuer, receiver) and source_kind == 'xml':
        raise DocumentError(f'{location}: XML no corresponde al RUC de esta empresa')
    if receiver and receiver != ruc and source_kind == 'registro':
        raise DocumentError(f'{location}: Registro de compras de otro RUC')
    return dict(source_kind=source_kind, issuer_ruc=issuer, receiver_ruc=receiver,
                voucher_id=voucher, voucher_type=typ,
                issue_date=_date(mapping.get('issue_date')), currency=currency,
                total=_amount(mapping.get('total')), source_location=location)


def _text(node, path):
    match = node.find(path)
    return (match.text or '').strip() if match is not None and match.text else ''


def _ruc_of_party(root, which):
    el = root.find(CAC + which)
    if el is None:
        return ''
    return (_text(el, './' + CAC + 'Party/' + CAC + 'PartyIdentification/' + CBC + 'ID')
            or _text(el, './' + CBC + 'CustomerAssignedAccountID'))


def _parse_xml(data: bytes, ruc):
    try:
        root = SafeET.fromstring(data, forbid_dtd=True, forbid_entities=True, forbid_external=True)
    except Exception as exc:
        raise DocumentError(f'XML rechazado por parseo seguro: {exc}') from exc
    if root.tag != '{urn:oasis:names:specification:ubl:schema:xsd:Invoice-2}Invoice':
        raise DocumentError('XML distinto de factura UBL Invoice (notas y boletas: fuera de este adaptador)')
    info = dict(issuer_ruc=_ruc_of_party(root, 'AccountingSupplierParty'),
                receiver_ruc=_ruc_of_party(root, 'AccountingCustomerParty'),
                voucher_id=_text(root, './' + CBC + 'ID'),
                issue_date=_text(root, './' + CBC + 'IssueDate'),
                currency=_text(root, './' + CBC + 'DocumentCurrencyCode'),
                voucher_type=_text(root, './' + CBC + 'InvoiceTypeCode') or '01',
                total=_text(root, './' + CAC + 'LegalMonetaryTotal/' + CBC + 'PayableAmount'))
    rec = _record(info, 'xml', '/Invoice', ruc)
    return [rec], 'Factura UBL interpretada localmente; aceptación SUNAT no verificada.'


def _headers(head):
    values = {_normalized(v): i for i,v in enumerate(head) if v is not None}
    selection = {key: next((values[v] for v in options if v in values), None)
                 for key, options in FIELDS.items()}
    if any(selection[k] is None for k in REQUIRED):
        return None
    return selection


def _tabular(rows, sheet, ruc):
    try:
        header = next(rows)
    except StopIteration:
        return [], 'Hoja vacía; no hay registros.'
    columns = _headers(header)
    if not columns:
        return [], 'Sin columnas de registro de compras reconocidas; solo se indexó el documento.'
    result = []
    for number, row in enumerate(rows, start=2):
        if number > MAX_ROWS + 1:
            raise DocumentError('Superado límite de 50 000 filas')
        if not any(x is not None and str(x).strip() for x in row):
            continue
        if len(row) > len(header) and any(x is not None and str(x).strip() for x in row[len(header):]):
            raise DocumentError(f'{sheet}:fila {number}: más columnas que la cabecera; verificar separadores decimales/CSV')
        data = {key: row[pos] if pos is not None and pos < len(row) else ''
                for key, pos in columns.items()}
        location = f'{sheet}:fila {number}'
        result.append(_record(data, 'registro', location, ruc))
    return result, f'{len(result)} registros de compras reconocidos.'


def _parse_csv(data, ruc):
    decoded = None
    for encoding in ('utf-8-sig', 'cp1252'):
        try:
            decoded = data.decode(encoding)
            break
        except UnicodeDecodeError:
            pass
    if decoded is None:
        raise DocumentError('Codificación CSV incompatible (UTF-8 o CP1252)')
    try:
        dialect = csv.Sniffer().sniff(decoded[:8192], delimiters=',;\t|')
    except csv.Error:
        dialect = csv.excel
    return _tabular(iter(csv.reader(io.StringIO(decoded), dialect)), 'CSV', ruc)


def _parse_xlsx(path, ruc):
    try:
        with zipfile.ZipFile(path) as archive:
            if sum(x.file_size for x in archive.infolist()) > MAX_UNPACKED:
                raise DocumentError('XLSX excede el límite de expansión del ZIP')
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        try:
            all_records = []
            notes = []
            if len(workbook.worksheets) > 50:
                raise DocumentError('Máximo 50 hojas por libro.')
            for sheet in workbook:
                if sheet.max_column > 100 or sheet.max_row > MAX_ROWS + 1:
                    raise DocumentError(f'{sheet.title}: hoja demasiado grande (máximo 100 columnas y 50 000 filas).')
                records, note = _tabular(iter(sheet.values), sheet.title, ruc)
                all_records.extend(records)
                notes.append(f'{sheet.title}: {note}')
            return all_records, ' | '.join(notes)
        finally:
            workbook.close()
    except (zipfile.BadZipFile, OSError) as exc:
        raise DocumentError(f'XLSX inválido: {exc}') from exc


def _parse_pdf(path, ruc):
    try:
        from pypdf import PdfReader
        reader = PdfReader(path, strict=True)
        n=len(reader.pages)
        texts=[(reader.pages[i].extract_text() or '') for i in range(min(n,30))]
        if not any(s.strip() for s in texts):
            return [], 'PDF sin texto digital legible en las primeras 30 páginas. OCR fuera de la PoC.'
        return [], f'PDF indexado como evidencia ({n} páginas), sin extraer registros contables.'
    except Exception as exc:
        raise DocumentError(f'PDF inválido: {exc}') from exc


def read_document(path, ruc):
    ext = path.suffix.lower()
    if ext not in FORMATS:
        raise DocumentError(f'Extensión no soportada: {ext}')
    if not path.is_file() or path.stat().st_size > MAX_SIZE:
        raise DocumentError('Archivo no accesible o supera 20 MiB')
    if ext == '.xml':
        return _parse_xml(path.read_bytes(), ruc)
    if ext == '.csv':
        return _parse_csv(path.read_bytes(), ruc)
    if ext == '.xlsx':
        return _parse_xlsx(path, ruc)
    return _parse_pdf(path, ruc)


def scan(conn, company, max_files=1000):
    """No se sigue symlink fuera de raíces permitidas. Originales nunca modificados."""
    cid = company['id']
    ruc = company['ruc']
    scanned, new, unchanged, duplicate, errors = 0,0,0,0,0
    notes = []
    found = set()
    scanned_roots = []
    for root in roots(conn, cid):
        if not root.is_dir():
            notes.append(f'Carpeta inaccesible: {root}')
            errors += 1
            continue
        scanned_roots.append(root.resolve())
        for path in sorted(root.rglob('*')):
            if path.suffix.lower() not in FORMATS or not path.is_file():
                continue
            resolved = path.resolve()
            if not resolved.is_relative_to(root.resolve()) or resolved in found:
                continue
            found.add(resolved)
            scanned += 1
            if scanned > max_files:
                raise DocumentError('Límite de archivos por escaneo alcanzado')
            kind = FORMATS[path.suffix.lower()]
            try:
                if path.stat().st_size > MAX_SIZE:
                    raise DocumentError('Archivo supera 20 MiB')
                with path.open('rb') as stream:
                    digest = hashlib.file_digest(stream, 'sha256').hexdigest()
            except OSError as exc:
                notes.append(f'{path.name}: {exc}')
                errors += 1
                continue
            existing = conn.execute('SELECT id,sha256,status FROM documents WHERE company_id=? AND path=?', (cid,str(resolved))).fetchone()
            if existing and existing['sha256'] == digest:
                # Un duplicado debe convertirse en fuente activa si su original desaparece.
                promote = (existing['status'] == 'duplicate' and not conn.execute(
                    "SELECT 1 FROM documents WHERE company_id=? AND sha256=? AND status='ok' AND path<>? LIMIT 1",
                    (cid,digest,str(resolved))).fetchone())
                if not promote:
                    unchanged += 1
                    continue
            elsewhere = conn.execute("SELECT id FROM documents WHERE company_id=? AND sha256=? AND path<>? AND status='ok' LIMIT 1",(cid,digest,str(resolved))).fetchone()
            if elsewhere:
                recs, status, note = [], 'duplicate', 'Mismo contenido que otro archivo del mismo RUC; no se duplica el conteo.'
                duplicate += 1
            else:
                try:
                    before = resolved.stat()
                    recs, note = read_document(resolved, ruc)
                    after = resolved.stat()
                    if before.st_size != after.st_size or before.st_mtime_ns != after.st_mtime_ns:
                        raise DocumentError('Archivo cambiado durante la lectura; volver a escanear.')
                    status = 'ok' if recs else 'indexed_only'
                    new += 1
                except Exception as exc:
                    recs, status, note = [], 'error', str(exc)
                    errors += 1
            with conn:
                if existing:
                    docid = existing['id']
                    conn.execute('DELETE FROM records WHERE document_id=? AND company_id=?', (docid,cid))
                    conn.execute('UPDATE documents SET sha256=?,kind=?,status=?,note=?,imported_at=CURRENT_TIMESTAMP WHERE id=? AND company_id=?',
                                 (digest,kind,status,note,docid,cid))
                else:
                    docid = conn.execute('INSERT INTO documents(company_id,path,sha256,kind,status,note) VALUES(?,?,?,?,?,?)',
                                         (cid,str(resolved),digest,kind,status,note)).lastrowid
                conn.executemany('''INSERT INTO records (company_id,document_id,source_kind,issuer_ruc,receiver_ruc,voucher_id,voucher_type,issue_date,currency,total,source_location)
                                    VALUES (:company_id,:document_id,:source_kind,:issuer_ruc,:receiver_ruc,:voucher_id,:voucher_type,:issue_date,:currency,:total,:source_location)''',
                                 [dict(r, company_id=cid, document_id=docid) for r in recs])
            notes.append(f'{path.name}: {status} - {note}')
    # Evitar informes falsos al conservar facturas de archivos eliminados.
    missing = 0
    for doc in conn.execute('SELECT id,path,status FROM documents WHERE company_id=?',(cid,)).fetchall():
        path = Path(doc['path'])
        if any(path.is_relative_to(root) for root in scanned_roots) and path not in found and doc['status'] != 'missing':
            with conn:
                conn.execute('DELETE FROM records WHERE company_id=? AND document_id=?',(cid,doc['id']))
                conn.execute("UPDATE documents SET status='missing',note='El archivo ya no existe en la carpeta autorizada.' WHERE company_id=? AND id=?",(cid,doc['id']))
            missing += 1
            notes.append(f'{path.name}: missing - eliminado del cómputo')
    return dict(scanned=scanned, imported=new, unchanged=unchanged, duplicates=duplicate, missing=missing, errors=errors, notes=notes)
